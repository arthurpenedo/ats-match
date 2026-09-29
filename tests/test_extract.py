from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from ats_match.cli import main
from ats_match.extract import ExtractionError, extract_text


def text_pdf(lines: list[str]) -> bytes:
    """Gera um PDF mínimo com texto real (sem dependências extras)."""
    content = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(f"({line}) '" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


LINES = ["Analista de dados", "Automatizei relatorios com Python e SQL no AWS Athena.",
         "Criei dashboards em Power BI para a lideranca de atendimento.",
         "Participei de testes de prompts de um chatbot com IA generativa."]


def test_pdf_with_text_is_extracted():
    text = extract_text(text_pdf(LINES), "cv.pdf")
    assert "Python e SQL" in text and "Power BI" in text


def test_image_only_pdf_is_rejected_with_explanation():
    buf = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(buf)
    with pytest.raises(ExtractionError, match="parece ser uma imagem"):
        extract_text(buf.getvalue(), "escaneado.pdf")


def test_docx_reads_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Experiência com Python")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "SQL"
    table.rows[0].cells[1].text = "Power BI"
    buf = BytesIO()
    doc.save(buf)
    text = extract_text(buf.getvalue(), "cv.docx")
    assert "Python" in text and "SQL" in text and "Power BI" in text


def test_unsupported_format():
    with pytest.raises(ExtractionError, match="não suportado"):
        extract_text(b"x", "cv.odt")


def test_cli_accepts_pdf(tmp_path, capsys):
    cv = tmp_path / "cv.pdf"
    cv.write_bytes(text_pdf(LINES))
    job = tmp_path / "vaga.md"
    job.write_text("## Requisitos\n- Python\n- SQL\n", encoding="utf-8")
    assert main([str(cv), str(job)]) == 0
    assert "Nota de aderência" in capsys.readouterr().out
